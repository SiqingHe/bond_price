import pandas as pd
from pathlib import Path
from openpyxl import load_workbook,Workbook
import time
from tqdm import tqdm
import csv

def sets_split(bond_pd):
    # sets=["train","valid","test"]
    split_dic={}
    bond_pd.fillna(-1,inplace=True)
    pd_len=bond_pd.shape[0]
    
    if pd_len==1:
        split_dic["train"]=bond_pd.copy()
    elif pd_len==2:
        split_dic["train"]=bond_pd.iloc[0:1]
        split_dic["valid"]=bond_pd.iloc[1:2]
    else:
        valid_num,test_num=int(pd_len/10)+1,int(pd_len/10)+1
        train_num=pd_len-valid_num-test_num
        split_dic["train"]=bond_pd.iloc[0:train_num]
        split_dic["valid"]=bond_pd.iloc[train_num:train_num+valid_num]
        split_dic["test"]=bond_pd.iloc[train_num+valid_num:train_num+valid_num+test_num]
    return split_dic

def set2_split():
    pass

def set3_split():
    pass

def data_add(table_dir,save_dir):
    Path(save_dir).mkdir(exist_ok=True,parents=True)
    # sheet_dic,sheet_dic_workbook={},{}
    csv_open,csv_write={},{}
    num=0
    for table_path in tqdm(Path(table_dir).glob("*.csv")):
        # if num>4:break
        tic=time.time()
        table_pd=pd.read_csv(table_path)
        table_pd=table_pd.loc[:,~table_pd.columns.str.contains("^Unnamed")]
        # print("*"*20)
        # print("total",table_pd.shape[0])
        split_dic=sets_split(table_pd)
        tbnum = 0
        for k,v in split_dic.items():
            # print(k,v.shape[0])
            tbnum += v.shape[0]
            save_path=str(Path(save_dir).joinpath(k+".csv"))
            if not Path(save_path).exists():
                with open(save_path,"w") as csv_open[k]:
                    csv_write[k]=csv.writer(csv_open[k])
                    csv_write[k].writerow(table_pd.columns.to_list())
            data = v.values.tolist()
            with open(save_path,"a") as csv_open[k]:
                csv_write[k]=csv.writer(csv_open[k])
                if not isinstance(data[0],list):
                    csv_write[k].writerow(data)
                else:
                    csv_write[k].writerows(v.values.tolist())
                # for row in data:
                #     csv_write[k].writerow(row[1:])
            # if k not in sheet_dic:
            #     sheet_dic_workbook[k] = Workbook()
            #     sheet_dic[k] = sheet_dic_workbook[k].active
            #     sheet_dic[k].append(table_pd.columns.to_list())
            # # 填充数据
            # data = v.values.tolist()
            # if not isinstance(data[0],list):
            #     sheet_dic[k].append(data)
            # else:
            #     for row in data:
            #         sheet_dic[k].append(row)
        if tbnum != table_pd.shape[0]:
            print("*"*20)
            print(table_path.name)
            print(table_pd.shape[0])
            for kk,vv in split_dic.items():
                print(vv.shapes[0])
        
        num+=1
        # print(time.time()-tic)
    # for k in sheet_dic.keys():
    #     sheet_dic_workbook[k].save(str(Path(save_dir).joinpath(k+".xlsx")))
        



if __name__=="__main__":
    pass
    # test_table=r"D:\python_code\LSTM-master\bond_price\real_data\test\010221.IB_41_test1.csv"
    # test_pd=pd.read_csv(test_table,index_col="org_date")
    # xy_split(test_pd,X_column,y_column)
    data_add(table_dir=r"D:\python_code\LSTM-master\bond_price\dealed_dir\dealed_0808",
             save_dir=r"D:\python_code\LSTM-master\bond_price\dealed_dir\sets_split0808")